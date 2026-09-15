import os

from docx import Document

from ..models import ParsedDocument


def parse_docx(file_path: str) -> ParsedDocument:
    """Extract text from a DOCX file, including any tables."""

    doc = Document(file_path)

    lines: list = []

    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            lines.append(paragraph.text)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    lines.append(cell.text)

    return ParsedDocument(
        source_filename=os.path.basename(file_path),
        raw_text="\n".join(lines),
        page_map=None,
        metadata={"paragraph_count": len(doc.paragraphs)},
        parse_status="success",
    )