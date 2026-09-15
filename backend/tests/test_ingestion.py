import os

import pymupdf
import pytest
from docx import Document as DocxDocument

from backend.app.ingestion.batch import batch_summary, process_batch
from backend.app.ingestion.dispatcher import parse_resume


@pytest.fixture
def sample_pdf(tmp_path):
    path = tmp_path / "sample.pdf"

    doc = pymupdf.open()

    page1 = doc.new_page()
    page1.insert_text(
        (72, 72),
        "Page one content: Python and FastAPI experience.",
    )

    page2 = doc.new_page()
    page2.insert_text(
        (72, 72),
        "Page two content: worked with PostgreSQL.",
    )

    doc.save(str(path))
    doc.close()

    return path


@pytest.fixture
def empty_pdf(tmp_path):
    """A blank PDF page with no extractable text."""

    path = tmp_path / "empty.pdf"

    doc = pymupdf.open()
    doc.new_page()
    doc.save(str(path))
    doc.close()

    return path


@pytest.fixture
def corrupted_pdf(tmp_path):
    path = tmp_path / "corrupted.pdf"
    path.write_bytes(b"this is not a real pdf file at all")

    return path


@pytest.fixture
def sample_docx(tmp_path):
    path = tmp_path / "sample.docx"

    doc = DocxDocument()

    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("SKILLS")
    doc.add_paragraph("Python, FastAPI, PostgreSQL")

    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Company"
    table.rows[0].cells[1].text = "Acme Corp"

    doc.save(str(path))

    return path


@pytest.fixture
def empty_docx(tmp_path):
    path = tmp_path / "empty.docx"

    doc = DocxDocument()
    doc.save(str(path))

    return path


@pytest.fixture
def table_only_docx(tmp_path):
    path = tmp_path / "table_only.docx"

    doc = DocxDocument()

    table = doc.add_table(rows=2, cols=2)

    table.rows[0].cells[0].text = "Skill"
    table.rows[0].cells[1].text = "Years"

    table.rows[1].cells[0].text = "Python"
    table.rows[1].cells[1].text = "3"

    doc.save(str(path))

    return path


@pytest.fixture
def sample_txt(tmp_path):
    path = tmp_path / "sample.txt"

    path.write_text(
        "Plain text resume content.\nSkills: Python, SQL.",
        encoding="utf-8",
    )

    return path


@pytest.fixture
def empty_txt(tmp_path):
    path = tmp_path / "empty.txt"

    path.write_text("", encoding="utf-8")

    return path


@pytest.fixture
def unsupported_file(tmp_path):
    path = tmp_path / "resume.xyz"

    path.write_text(
        "some content",
        encoding="utf-8",
    )

    return path


def test_pdf_text_extraction_and_page_mapping(sample_pdf):
    result = parse_resume(str(sample_pdf))

    assert result.parse_status == "success"
    assert "Python" in result.raw_text
    assert "PostgreSQL" in result.raw_text

    assert result.page_map is not None
    assert len(result.page_map) == 2

    assert "Python" in result.page_map[1]
    assert "PostgreSQL" in result.page_map[2]


def test_docx_paragraph_and_table_extraction(sample_docx):
    result = parse_resume(str(sample_docx))

    assert result.parse_status == "success"
    assert "Jane Doe" in result.raw_text
    assert "SKILLS" in result.raw_text
    assert "Acme Corp" in result.raw_text

    assert result.page_map is None
    assert result.metadata["paragraph_count"] == 3
    assert result.metadata["table_count"] == 1


def test_docx_with_only_a_table_extracts_table_content(
    table_only_docx,
):
    result = parse_resume(str(table_only_docx))

    assert result.parse_status == "success"
    assert "Python" in result.raw_text
    assert "3" in result.raw_text


def test_txt_extraction(sample_txt):
    result = parse_resume(str(sample_txt))

    assert result.parse_status == "success"
    assert "Plain text resume content." in result.raw_text
    assert result.page_map is None


def test_unsupported_file_type(unsupported_file):
    result = parse_resume(str(unsupported_file))

    assert result.parse_status == "failed"
    assert result.error is not None
    assert "Unsupported file type" in result.error


def test_corrupted_pdf_does_not_raise(corrupted_pdf):
    result = parse_resume(str(corrupted_pdf))

    assert result.parse_status == "failed"
    assert result.error is not None


def test_missing_file_does_not_raise(tmp_path):
    missing_file = tmp_path / "missing.pdf"

    result = parse_resume(str(missing_file))

    assert result.parse_status == "failed"
    assert result.error is not None


def test_empty_txt_is_marked_failed(empty_txt):
    result = parse_resume(str(empty_txt))

    assert result.parse_status == "failed"
    assert result.error is not None
    assert "No extractable text" in result.error


def test_empty_docx_is_marked_failed(empty_docx):
    result = parse_resume(str(empty_docx))

    assert result.parse_status == "failed"
    assert result.error is not None
    assert "No extractable text" in result.error


def test_blank_pdf_is_marked_failed(empty_pdf):
    result = parse_resume(str(empty_pdf))

    assert result.parse_status == "failed"
    assert result.error is not None
    assert "No extractable text" in result.error


def test_source_filename_is_basename_not_full_path(sample_txt):
    result = parse_resume(str(sample_txt))

    assert result.source_filename == "sample.txt"
    assert os.sep not in result.source_filename


def test_batch_processing_isolates_failures(
    sample_pdf,
    corrupted_pdf,
    sample_txt,
):
    results = process_batch(
        [
            str(sample_pdf),
            str(corrupted_pdf),
            str(sample_txt),
        ]
    )

    assert len(results) == 3

    statuses = [
        result.parse_status
        for result in results
    ]

    assert statuses == [
        "success",
        "failed",
        "success",
    ]


def test_batch_summary_counts(
    sample_pdf,
    corrupted_pdf,
    sample_txt,
    empty_txt,
):
    results = process_batch(
        [
            str(sample_pdf),
            str(corrupted_pdf),
            str(sample_txt),
            str(empty_txt),
        ]
    )

    summary = batch_summary(results)

    assert summary["total"] == 4
    assert summary["succeeded"] == 2
    assert summary["failed"] == 2
    assert len(summary["failed_files"]) == 2


def test_empty_batch():
    results = process_batch([])

    assert results == []

    summary = batch_summary(results)

    assert summary == {
        "total": 0,
        "succeeded": 0,
        "failed": 0,
        "failed_files": [],
    }